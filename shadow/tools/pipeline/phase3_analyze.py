#!/usr/bin/env python3
"""Phase 3 - statistical sheets. Reads <out>/<run>/<area>/phase2/, writes .../phase3/.

Per experiment: <run>__<area>.xlsx with the tabs wPoA_epoch, wPoA_longitudinal,
WeightEngine_epoch, config, manifest, notes, plus one CSV per table (same
content, diff-able). Every tab starts with the parameter row of the experiment.
Every metric that cannot be computed on this campaign carries the string
"not computable: <reason>" instead of a number.

Campaign level (--campaign): campaign_summary.csv / campaign_families.csv /
campaign_manifest.json / campaign.xlsx, and the standalone Monte Carlo validator
of the sortition algorithm (tools/valida_sortition_montecarlo.py) is invoked and
its output copied next to them.

    python3 tools/pipeline/phase3_analyze.py --out risultati [--only run2] [--campaign]
"""

import argparse
import json
import logging
import math
import statistics
import subprocess
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pipeline import common as C  # noqa: E402
from pipeline.stat import concentration as CONC  # noqa: E402
from pipeline.stat import gof as GOF  # noqa: E402
from pipeline.stat import longitudinal as LONG  # noqa: E402
from pipeline.stat import streak as STREAK  # noqa: E402
from pipeline.stat import timer_race as TR  # noqa: E402
from pipeline.stat import wilson as WIL  # noqa: E402

try:
    import openpyxl
    from openpyxl.styles import Font
    USING_OPENPYXL = True
except ImportError:                                   # pragma: no cover
    openpyxl = None
    USING_OPENPYXL = False

LOG = logging.getLogger("pipeline.phase3")

ALPHA = 0.05
MC_GOF = 20000
MC_STREAK = 10000
SEED = 20260905
NC = "not computable: "


def _r(x, nd=6):
    if isinstance(x, float):
        if math.isnan(x) or math.isinf(x):
            return ""
        return round(x, nd)
    return x


def _round_rows(rows, nd=6):
    return [OrderedDict((k, _r(v, nd)) for k, v in r.items()) for r in rows]


def param_header(man):
    p = man["params"]
    su = man.get("startup_params", {})
    return OrderedDict([
        ("experiment_id", man["experiment_id"]), ("area", man["area"]),
        ("replication_id", man["replication_id"]),
        ("target_block_time", man["target_block_time"]), ("weight_epoch_length", man["weight_epoch_length"]),
        ("setup_blocks", man["setup_blocks"]), ("final_height", man["final_height"]),
        ("wpoa_sortition_delta", p.get("wpoa_sortition_delta", "")),
        ("wpoa_sortition_lambda", p.get("wpoa_sortition_lambda", "")),
        ("wpoa_randao_lookback", p.get("wpoa_randao_lookback", "")),
        ("dumpfunction", man["dumpfunction"]), ("malus_enabled", man["malus_enabled"]),
        ("wpoa_malus_mu", p.get("wpoa_malus_mu", "")), ("wpoa_malus_max", p.get("wpoa_malus_max", "")),
        ("malus_points_equiv_delay_selfwrite_badweight", "%s/%s/%s/%s" % (
            p.get("wpoa_malus_equiv_points", ""), p.get("wpoa_malus_delay_points", ""),
            su.get("malus_selfwrite_points", p.get("wpoa_malus_selfwrite_points", "")),
            su.get("malus_badweight_points", p.get("wpoa_malus_badweight_points", "")))),
        ("weight_kappa", p.get("weight_kappa", "")), ("weight_alpha", p.get("weight_alpha", "")),
        ("weight_lambda", p.get("weight_lambda", "")), ("mining_diversity", p.get("mining_diversity", "")),
        ("shadow_seed", man.get("shadow", {}).get("general_seed", "")),
        ("poesia_rng_seed", man.get("shadow", {}).get("poesia_rng_seed", "")),
        ("rtt_max_between_miners_ms", man.get("shadow", {}).get("rtt_max_between_miners_ms", "")),
    ])


# --------------------------------------------------------------------------
# loading
# --------------------------------------------------------------------------

def load_phase2(p2):
    man = json.loads((p2 / "manifest.json").read_text(encoding="utf-8"))
    t = {n: C.read_csv_dicts(p2 / (n + ".csv")) for n in (
        "candidates_long", "round_level", "epoch_level_wpoa", "epoch_level_weight_engine",
        "weight_engine_verification", "malus_epoch")}
    return man, t


class RunData:
    """Typed views of the phase-2 tables."""

    def __init__(self, man, t):
        self.man = man
        self.tbt = float(man["target_block_time"])
        self.L = int(man["weight_epoch_length"])
        self.delta = C.fnum(man["params"].get("wpoa_sortition_delta"), 0.5)
        self.lam = C.fnum(man["params"].get("wpoa_sortition_lambda"), 0.0)
        self.dmax = self.delta * self.tbt
        self.dump = man["dumpfunction"] or "none"
        self.rounds = sorted((r for r in t["round_level"]), key=lambda r: int(r["height"]))
        self.by_epoch = defaultdict(list)
        for r in self.rounds:
            self.by_epoch[int(r["epoch"])].append(r)
        self.validators = sorted({r["validator_address"] for r in t["epoch_level_wpoa"]})
        self.host = {r["validator_address"]: r["validator_host"] for r in t["epoch_level_wpoa"]}
        self.epoch_rows = defaultdict(dict)            # epoch -> addr -> row
        for r in t["epoch_level_wpoa"]:
            self.epoch_rows[int(r["epoch"])][r["validator_address"]] = r
        self.cands = defaultdict(dict)                 # height -> addr -> row
        for r in t["candidates_long"]:
            self.cands[int(r["height"])][r["candidate_address"]] = r
        self.we = t["epoch_level_weight_engine"]
        self.ver = t["weight_engine_verification"]
        self.malus = t["malus_epoch"]
        ep = {e["epoch"]: e for e in man.get("epochs", [])}
        self.fully_measured = sorted(e for e in self.by_epoch if ep.get(e, {}).get("fully_measured")
                                     and ep.get(e, {}).get("end_height", 0) <= int(man["final_height"]))
        self.measured = sorted(self.by_epoch)

    def share_matrix(self, rounds):
        """(R, n) theoretical shares and (R, n) effective weights, W_tot (R,), phi (R,)."""
        S, W, WT, PH = [], [], [], []
        for r in rounds:
            h = int(r["height"])
            cs = self.cands.get(h, {})
            row_s, row_w = [], []
            for a in self.validators:
                c = cs.get(a)
                row_s.append(C.fnum(c["p_theoretical"], 0.0) if c else 0.0)
                row_w.append(C.fnum(c["weight_effective"], 0.0) if c else 0.0)
            if sum(row_s) <= 0:
                continue
            S.append(row_s)
            W.append(row_w)
            WT.append(C.fnum(r["W_tot_effective"], sum(row_w)))
            PH.append(C.fnum(r["phi_s"], 0.0))
        return np.array(S), np.array(W), np.array(WT), np.array(PH)


# --------------------------------------------------------------------------
# sigma estimates (phase-0 report, section 5.6)
# --------------------------------------------------------------------------

def sigma_rows(rd, base):
    res = [C.fnum(r["scheduler_residual_s"]) for r in rd.rounds if r["scheduler_residual_s"] != ""]
    res = [x for x in res if x is not None]
    gaps = [C.fnum(r["inversion_gap_s"]) for r in rd.rounds if r["inversion"] == "1" and r["inversion_gap_s"] != ""]
    rtt_ms = C.fnum(rd.man.get("shadow", {}).get("rtt_max_between_miners_ms"))
    rows = []
    s1 = (rtt_ms / 2.0 / 1000.0) if rtt_ms else None
    rows.append(dict(base, sigma_id="S1_topology", sigma_s=s1 if s1 is not None else "",
                     source="max one-way shortest-path latency between miners from the .gml (rtt/2)",
                     what_it_measures="propagation latency; deterministic (jitter 0 us), not a variance",
                     n=len(rd.man.get("miners", [])), median_s="", p90_s="", max_s="",
                     note="" if s1 is not None else NC + "topology matrix absent"))
    if len(res) >= 2:
        rows.append(dict(base, sigma_id="S2_scheduler_residual_sd", sigma_s=statistics.pstdev(res),
                         source="sd of block_time - parent_time - D_logged(observed winner) over measured rounds",
                         what_it_measures="quantisation of the 1 s time bar and block timestamps plus miner loop granularity: the noise that actually decides the race",
                         n=len(res), median_s=statistics.median(res), p90_s=sorted(res)[int(0.9 * (len(res) - 1))],
                         max_s=max(res), note="PRIMARY estimate for the Prop. 5.18 bound"))
    else:
        rows.append(dict(base, sigma_id="S2_scheduler_residual_sd", sigma_s="", source="", what_it_measures="",
                         n=len(res), median_s="", p90_s="", max_s="", note=NC + "fewer than 2 residuals"))
    if gaps:
        rows.append(dict(base, sigma_id="S2b_inversion_gap", sigma_s=statistics.pstdev(gaps) if len(gaps) > 1 else "",
                         source="D_observed_winner - D_min on inverted rounds",
                         what_it_measures="size of the disadvantage overcome by the observed winner (only on inverted rounds; biased by selection)",
                         n=len(gaps), median_s=statistics.median(gaps), p90_s=sorted(gaps)[int(0.9 * (len(gaps) - 1))],
                         max_s=max(gaps), note="reported for comparison, not used in the bound"))
    rows.append(dict(base, sigma_id="S3_packet_loss", sigma_s="", source="", what_it_measures="",
                     n="", median_s="", p90_s="", max_s="", note=NC + "packet loss is not observable from the archived logs (1 s resolution, no per-message trace)"))
    return rows, (s1, statistics.pstdev(res) if len(res) >= 2 else None)


# --------------------------------------------------------------------------
# wPoA per epoch
# --------------------------------------------------------------------------

def wpoa_epoch(rd, base, sigmas):
    s1, s2 = sigmas
    val_rows, test_rows, live_rows, p517_rows = [], [], [], []
    groups = [(e, rd.by_epoch[e]) for e in rd.measured] + [("all", rd.rounds)]
    nval = len(rd.validators)
    for e, rounds in groups:
        B = len(rounds)
        S, W, WT, PH = rd.share_matrix(rounds)
        obs = Counter(r["winner_observed_address"] for r in rounds)
        O = [obs.get(a, 0) for a in rd.validators]
        E = list(S.sum(axis=0)) if len(S) else [0.0] * nval
        p_bar = [x / len(S) if len(S) else None for x in E]
        p_hat = [o / B if B else None for o in O]
        buried = int(all(r["epoch_buried"] == "1" for r in rounds)) if rounds else ""
        # --- validators table (weights, p_i, Wilson) ---
        for i, a in enumerate(rd.validators):
            er = rd.epoch_rows.get(e, {}).get(a) if e != "all" else None
            lo, hi = WIL.wilson_interval(O[i], B)
            row = OrderedDict(base)
            row.update({
                "epoch": e, "epoch_buried": buried, "validator_host": rd.host.get(a, ""), "validator_address": a,
                "w_raw_start": er["w_raw_start"] if er else "", "w_raw_end": er["w_raw_end"] if er else "",
                "w_eff_start": er["w_eff_start"] if er else "", "w_eff_end": er["w_eff_end"] if er else "",
                "psi": 1.0 if rd.malus else "", "dumpfunction": rd.dump,
                "p_theoretical_blockweighted": p_bar[i] if p_bar[i] is not None else "",
                "n_blocks": B, "O_i": O[i], "E_i": E[i],
                "p_hat": p_hat[i] if p_hat[i] is not None else "",
                "wilson95_low": lo if lo is not None else "", "wilson95_high": hi if hi is not None else "",
                "p_theoretical_inside_wilson95": int(lo <= p_bar[i] <= hi) if (lo is not None and p_bar[i] is not None) else "",
                "abs_error_p": abs(p_hat[i] - p_bar[i]) if (p_hat[i] is not None and p_bar[i] is not None) else "",
                "share_changed_within_epoch": er["share_changed_within_epoch"] if er else "",
            })
            val_rows.append(row)
            # liveness per validator: MissedTurnRate
            if e != "all":
                am, lost = C.inum(er["n_rounds_theoretical_argmin"], 0), C.inum(er["n_rounds_theoretical_argmin_lost"], 0)
            else:
                am = sum(1 for r in rounds if r["winner_theoretical_argmin_delay_host"] == rd.host.get(a))
                lost = sum(1 for r in rounds if r["winner_theoretical_argmin_delay_host"] == rd.host.get(a) and r["inversion"] == "1")
            lrow = OrderedDict(base)
            lrow.update({"epoch": e, "validator_host": rd.host.get(a, ""), "validator_address": a,
                         "n_rounds_theoretical_argmin": am, "n_rounds_theoretical_argmin_lost": lost,
                         "MissedTurnRate": (lost / am) if am else "",
                         "MissedTurnRate_note": "" if am else NC + "never the theoretical argmin in this epoch",
                         "n_rounds_won_not_argmin": sum(1 for r in rounds if r["winner_observed_address"] == a and r["inversion"] == "1"),
                         "blocks_won": O[i]})
            live_rows.append(lrow)
        # --- tests table ---
        tr = OrderedDict(base)
        tr.update({"epoch": e, "epoch_buried": buried, "n_blocks": B, "n_rounds_with_scores": len(S),
                   "n_validators": nval, "share_changes_within_epoch": int(bool(len(S)) and bool(np.any(np.ptp(S, axis=0) > 1e-12)))})
        # GoF
        g = GOF.gof_test(O, E, alpha=ALPHA, mc_draws=MC_GOF, seed=SEED)
        tr.update({"gof_method": g["gof_method"], "gof_chi2": g["chi2_stat"], "gof_df": g["df"],
                   "gof_min_expected": g["min_expected"], "gof_p_value": g["p_value"],
                   "gof_reject_alpha05": g["reject_alpha"], "gof_exact_outcomes": g["exact_outcomes_enumerated"],
                   "MAE_p": g.get("MAE_p", ""), "MaxAE_p": g.get("MaxAE_p", ""), "TV": g.get("TV", "")})
        if len(S) >= 1 and sum(O) > 0:
            pv, st, _ = GOF.mc_round_by_round_pvalue(O, S, draws=MC_GOF, seed=SEED)
            tr["gof_mc_round_by_round_p"] = pv
            tr["gof_mc_round_by_round_draws"] = MC_GOF
        else:
            tr["gof_mc_round_by_round_p"] = NC + "no rounds with scores"
            tr["gof_mc_round_by_round_draws"] = ""
        # timer race
        margins = [C.fnum(r["margin_G_logged_s"]) for r in rounds if r["margin_G_logged_s"] != ""]
        ncand = Counter(C.inum(r["n_candidates_with_weight"], 0) for r in rounds).most_common(1)
        n_c = ncand[0][0] if ncand else nval
        kb = TR.ks_beta_margin(margins, rd.dmax, n_c)
        tr.update({"G_n": len(margins), "G_mean_s": statistics.mean(margins) if margins else "",
                   "G_median_s": statistics.median(margins) if margins else "",
                   "G_over_Dmax_median": (statistics.median(margins) / rd.dmax) if (margins and rd.dmax) else "",
                   "ks_beta_stat": kb["ks_beta_stat"], "ks_beta_p": kb["ks_beta_p"],
                   "ks_beta_mean_G_s": kb.get("beta_mean_G_s", ""), "ks_beta_note": kb["ks_beta_note"]})
        km = TR.ks_exact_mc_margin(margins, W, WT, rd.tbt, rd.delta, rd.lam, PH, draws=MC_GOF, seed=SEED) \
            if len(S) else {"ks_mc_stat": "", "ks_mc_p": "", "mc_draws": "", "ks_mc_note": NC + "no rounds"}
        tr.update({"ks_mc_stat": km["ks_mc_stat"], "ks_mc_p": km["ks_mc_p"], "ks_mc_mean_G_s": km.get("mc_mean_G_s", ""),
                   "ks_mc_draws": km["mc_draws"], "ks_mc_note": km["ks_mc_note"]})
        n_inv_rounds = [r for r in rounds if r["inversion"] != ""]
        n_inv = sum(1 for r in n_inv_rounds if r["inversion"] == "1")
        lo, hi = WIL.wilson_interval(n_inv, len(n_inv_rounds))
        tr.update({"inversion_n": n_inv, "inversion_rounds": len(n_inv_rounds),
                   "inversion_rate": (n_inv / len(n_inv_rounds)) if n_inv_rounds else "",
                   "inversion_wilson95_low": lo if lo is not None else "", "inversion_wilson95_high": hi if hi is not None else "",
                   "inversion_with_fork_n": sum(1 for r in n_inv_rounds if r["inversion"] == "1" and r["fork_detected"] == "1"),
                   "inversion_bound_sigma_S1": TR.inversion_bound(n_c, s1, rd.dmax) if s1 is not None else NC + "S1 absent",
                   "inversion_bound_sigma_S2": TR.inversion_bound(n_c, s2, rd.dmax) if s2 is not None else NC + "S2 absent",
                   "inversion_bound_sigma_used_s": s2 if s2 is not None else "",
                   "inversion_mc_prob_gaussian_sigma_S2": TR.mc_inversion_probability(W, WT, rd.tbt, rd.delta, rd.lam, PH, s2, draws=MC_GOF, seed=SEED) if (len(S) and s2) else "",
                   "inversion_bound_note": "bound (3/2)(n sigma/Dmax)^(2/3) of Prop. 5.18; sigma = S2 (scheduler) primary, S1 (topology) alongside; capped at 1"})
        # concentration
        tr.update(CONC.concentration_row([p for p in p_bar if p is not None], "p_theoretical_"))
        tr.update(CONC.concentration_row([p for p in p_hat if p is not None], "p_hat_"))
        # streaks
        seq = [r["winner_observed_address"] for r in rounds]
        so = STREAK.observed_streaks(seq)
        sm = STREAK.mc_streaks(S, so["L_max_observed"], so["repeat_prob_observed"], draws=MC_STREAK, seed=SEED) if len(S) >= 2 else {}
        tr.update({"L_max_observed": so["L_max_observed"], "repeat_prob_observed": so["repeat_prob_observed"],
                   "L_max_mc_mean": sm.get("L_max_mc_mean", ""), "L_max_mc_p95": sm.get("L_max_mc_p95", ""),
                   "L_max_mc_p_value": sm.get("L_max_mc_p_value", ""), "repeat_prob_mc_mean": sm.get("repeat_prob_mc_mean", ""),
                   "repeat_prob_mc_p_value": sm.get("repeat_prob_mc_p_value", ""), "streak_mc_draws": sm.get("mc_draws", "")})
        # liveness
        n_comp = sum(C.inum(r["n_competing_blocks"], 0) for r in rounds)
        dts = sorted(C.fnum(r["dt_prev_s"]) for r in rounds if r["dt_prev_s"] != "")
        def pct(q):
            return dts[min(len(dts) - 1, int(q * (len(dts) - 1)))] if dts else ""
        tr.update({"BlockSuccessRate": (B / (B + n_comp)) if (B + n_comp) else "",
                   "BlockSuccessRate_note": "canonical / (canonical + non-canonical competing blocks seen in `sortition OK` lines)",
                   "n_competing_blocks": n_comp, "n_forks": sum(1 for r in rounds if r["fork_detected"] == "1"),
                   "n_rejects": sum(C.inum(r["n_rejects"], 0) for r in rounds),
                   "rejects_by_reason": NC + "no REJECT line in the campaign" if not any(C.inum(r["n_rejects"], 0) for r in rounds) else "see phase1/rejects.csv",
                   "InterBlockTime_median_s": pct(0.5), "InterBlockTime_p95_s": pct(0.95), "InterBlockTime_p99_s": pct(0.99),
                   "InterBlockTime_mean_s": statistics.mean(dts) if dts else "", "InterBlockTime_max_s": dts[-1] if dts else "",
                   "time_bar_violations": sum(1 for r in rounds if r["time_bar_ok"] == "0"),
                   "time_bar_checked": sum(1 for r in rounds if r["time_bar_ok"] != ""),
                   "delay_recompute_mismatch_rounds": sum(1 for r in rounds if C.inum(r.get("n_candidates_delay_recompute_mismatch"), 0)),
                   "phi_mismatch_rounds": sum(1 for r in rounds if r["phi_identical_across_nodes"] == "0"),
                   "malus_delay_events_recomputed": NC + "malus registry empty (0 records): no Delay event to recompute; the time bar nTime >= parent + floor(D) was checked on every block instead",
                   "malus_psi": "1.0 for every validator (empty registry)"})
        test_rows.append(tr)
        # Prop. 5.17 per winner class (argmin by score) -- gaps standardised by (W - w_i*)
        gaps, wt, we = defaultdict(list), defaultdict(list), defaultdict(list)
        for r in rounds:
            if r["score_gap_2_1"] != "" and r["W_tot_minus_argmin_weff"] != "":
                host = r["winner_theoretical_argmin_score_host"]
                gaps[host].append(C.fnum(r["score_gap_2_1"]))
                wt[host].append(C.fnum(r["W_tot_effective"]))
                we[host].append(C.fnum(r["weff_argmin_score"]))
        for pr in TR.prop517_test(gaps, wt, we):
            row = OrderedDict(base)
            row.update({"epoch": e, "argmin_class_host": pr["winner_class"], "n": pr["n"],
                        "ks_stat": pr["ks_stat"], "ks_p": pr["ks_p"], "mean_gap_x_rate": pr.get("mean_gap_x_rate", ""),
                        "expected_mean": pr.get("expected_mean", ""), "note": pr["note"]})
            p517_rows.append(row)
    return val_rows, test_rows, live_rows, p517_rows


# --------------------------------------------------------------------------
# wPoA longitudinal
# --------------------------------------------------------------------------

def wpoa_longitudinal(rd, base):
    eps = rd.fully_measured or rd.measured
    if len(eps) < 2:
        note = NC + "fewer than 2 fully measured epochs"
        return ([dict(base, note=note)], [dict(base, note=note)], [], [dict(base, note=note)])
    e0, e1 = eps[0], eps[-1]

    def snap(e):
        out = {}
        for a in rd.validators:
            er = rd.epoch_rows[e].get(a)
            if er:
                out[a] = (C.inum(er["O_i"], 0), C.inum(er["n_blocks_epoch"], 0), C.fnum(er["p_theoretical_blockweighted"]))
        return out
    gain_rows = []
    for g in LONG.gains(snap(e0), snap(e1)):
        row = OrderedDict(base)
        row.update({"epoch_first": e0, "epoch_last": e1, "validator_host": rd.host.get(g["validator_address"], "")})
        row.update(g)
        # monotonicity over all measured epochs
        series = []
        for e in eps:
            er = rd.epoch_rows[e].get(g["validator_address"])
            if er and er["n_blocks_epoch"] not in ("", "0"):
                series.append((e, C.fnum(er["w_eff_end"]), C.inum(er["O_i"], 0) / C.inum(er["n_blocks_epoch"], 1)))
        m = LONG.monotonicity(series)
        row.update({"mono_" + k: v for k, v in m.items()})
        row["mono_note"] = "consecutive-epoch pairs: sign(dw_eff) vs sign(dp_hat), ties excluded; exact binomial sign test vs 1/2"
        gain_rows.append(row)

    # binomial logit GLM on (epoch x validator) points
    succ, trials, x_logw, x_logp, table = [], [], [], [], []
    for e in eps:
        for a in rd.validators:
            er = rd.epoch_rows[e].get(a)
            if not er or er["n_blocks_epoch"] in ("", "0") or er["w_eff_end"] == "" or C.fnum(er["p_theoretical_blockweighted"], 0) <= 0:
                continue
            succ.append(C.inum(er["O_i"], 0))
            trials.append(C.inum(er["n_blocks_epoch"], 0))
            x_logw.append(math.log(C.fnum(er["w_eff_end"])))
            x_logp.append(math.log(C.fnum(er["p_theoretical_blockweighted"])))
            table.append({"epoch": e, "address": a, "O": succ[-1], "B": trials[-1], "p": C.fnum(er["p_theoretical_blockweighted"])})
    fit_rows = []
    for name, x, h0 in (("logit(p_hat) ~ b0 + b1*log(w_eff)", x_logw, "b1 = 1 (proportional selection)"),
                        ("logit(p_hat) ~ b0 + b1*log(p_theoretical)", x_logp, "b1 = 1, b0 = 0")):
        f = LONG.binomial_logit_glm(succ, trials, x)
        row = OrderedDict(base)
        row.update({"model": name, "H0": h0})
        row.update(f)
        row["beta1_ci_contains_1"] = int(f["beta1_ci_low"] <= 1.0 <= f["beta1_ci_high"]) if f["beta1"] != "" else ""
        fit_rows.append(row)
    lr_rows = LONG.log_ratio_rows(table)
    lr_fit = LONG.log_ratio_fit(lr_rows)
    row = OrderedDict(base)
    row.update({"model": "log(p_hat_i/p_hat_j) ~ a + b*log(w_i/w_j) (pairs x epochs)", "H0": "b = 1, a = 0"})
    row.update({"beta1": lr_fit.get("slope", ""), "beta0": lr_fit.get("intercept", ""), "n_points": lr_fit.get("n_pairs", ""),
                "pearson_r": lr_fit.get("pearson_r", ""), "pearson_p": lr_fit.get("pearson_p", ""),
                "spearman_rho": lr_fit.get("spearman_rho", ""), "spearman_p": lr_fit.get("spearman_p", ""),
                "note": lr_fit.get("note", "")})
    fit_rows.append(row)
    lr_out = []
    for r in lr_rows:
        row = OrderedDict(base)
        row.update(r)
        row["validator_i_host"] = rd.host.get(r["validator_i"], "")
        row["validator_j_host"] = rd.host.get(r["validator_j"], "")
        lr_out.append(row)
    return gain_rows, fit_rows, lr_out, []


# --------------------------------------------------------------------------
# weight engine per epoch
# --------------------------------------------------------------------------

def weight_engine(rd, base):
    by_e = defaultdict(list)
    for r in rd.we:
        by_e[int(r["epoch"])].append(r)
    ver = {int(r["epoch"]): r for r in rd.ver}
    epoch_rows, gini_rows = [], []
    for e in sorted(by_e):
        rows = by_e[e]
        pubs = [C.fnum(r["w_k_published"]) for r in rows if r["w_k_published"] != ""]
        raws = [C.fnum(r["W_k_raw"]) for r in rows if r["W_k_raw"] != ""]
        cis = [C.fnum(r["c_i_sum"]) for r in rows]
        v = ver.get(e, {})
        grow = OrderedDict(base)
        grow.update({
            "epoch": e, "n_clusters": len(rows), "n_published": len(pubs),
            "gini_w_published": CONC.gini(pubs) if len(pubs) >= 2 else "",
            "gini_input_W_k_raw": CONC.gini(raws) if len(raws) >= 2 else "",
            "gini_input_c_i_sum": CONC.gini(cis) if len(cis) >= 2 else "",
            "gini_delta_published_minus_input": (CONC.gini(pubs) - CONC.gini(raws)) if (len(pubs) >= 2 and len(raws) >= 2) else "",
            "entropy_norm_w_published": CONC.normalized_entropy(pubs) if len(pubs) >= 2 else "",
            "entropy_norm_input_W_k_raw": CONC.normalized_entropy(raws) if len(raws) >= 2 else "",
            "HHI_w_published": CONC.hhi(pubs) if len(pubs) >= 2 else "",
            "dominant_host": max(rows, key=lambda r: C.fnum(r["w_k_published"], -1))["cluster_head_host"] if pubs else "",
            "dominant_share": (max(pubs) / sum(pubs)) if pubs and sum(pubs) > 0 else "",
            "n_recompute_match": sum(1 for r in rows if r["recompute_match"] == "1"),
            "n_recompute_mismatch": sum(1 for r in rows if r["recompute_match"] == "0"),
            "n_hosts_verification_ok_log": v.get("n_hosts_verification_ok", ""),
            "n_hosts_verification_failed_log": v.get("n_hosts_verification_failed", ""),
            "gini_note": "gini_input is the Gini of W_k^(e) (ESG x activity before the rho feedback): the baseline the published weight is compared with; if gini_published - gini_input grows over epochs the feedback amplifies inequality beyond the inputs",
        })
        gini_rows.append(grow)
        for r in rows:
            row = OrderedDict(base)
            row.update(r)
            row["dumpfunction_effect"] = (NC + "dump-function=none in this run: g(w) = w, no damping to measure") if rd.dump == "none" \
                else ("share with g: %s vs share without g: %s" % (r["share_snapshot"], ""))
            row["malus_trajectory"] = NC + "malus registry empty: M = 0 and Psi = 1 in every epoch (no decay to observe)"
            epoch_rows.append(row)

    # lag-1 Spearman rho_k(e) -> w_k(e+1), per cluster and pooled
    corr_rows = []
    pooled_x, pooled_y, pooled_yc, pooled_r = [], [], [], []
    clusters = sorted({r["cluster_head_address"] for r in rd.we})
    for m in clusters:
        rs = {int(r["epoch"]): r for r in rd.we if r["cluster_head_address"] == m}
        xs, ys, ycs, rks = [], [], [], []
        for e in sorted(rs):
            nxt = rs.get(e + 1)
            if not nxt or nxt["w_k_published"] == "" or rs[e]["rho_k"] == "":
                continue
            xs.append(C.fnum(rs[e]["rho_k"]))
            ys.append(C.fnum(nxt["w_k_published"]))
            raw_next = C.fnum(nxt["W_k_raw"])
            ycs.append((C.fnum(nxt["w_k_published"]) / raw_next) if raw_next else None)
            rks.append(C.fnum(rs[e]["R_k_reconciled"]))
        pooled_x += xs
        pooled_y += ys
        pooled_yc += ycs
        pooled_r += rks
        rho, p, n = C._spearman(xs, ys)
        rho_c, p_c, n_c = C._spearman(xs, ycs)
        rho_r, p_r, n_r = C._spearman(rks, xs)
        row = OrderedDict(base)
        row.update({"cluster_head_host": rd.host.get(m, m[:10]), "cluster_head_address": m, "n_epoch_pairs": n,
                    "n_distinct_rho_values": len(set(xs)), "n_pairs_rho_equal_1": sum(1 for x in xs if x is not None and abs(x - 1.0) < 1e-9),
                    "spearman_rho_e_vs_w_next": rho, "spearman_p": p,
                    "spearman_rho_e_vs_w_next_over_W_raw_next": rho_c, "spearman_p_control": p_c,
                    "spearman_R_k_vs_rho_k": rho_r, "spearman_p_R_vs_rho": p_r,
                    "note": ("" if n >= 3 else NC + "fewer than 3 epoch pairs; ") +
                            "the control against w(e+1)/W_k(e+1) = kappa[rho lambda_w + 1 - lambda_w] is deterministic by Def. 6.9 (tautological check); "
                            "with many ties at rho = 1 the coefficient is driven by the few epochs with incomplete reconciliation"})
        corr_rows.append(row)
    rho, p, n = C._spearman(pooled_x, pooled_y)
    rho_c, p_c, n_c = C._spearman(pooled_x, pooled_yc)
    rho_r, p_r, n_r = C._spearman(pooled_r, pooled_x)
    row = OrderedDict(base)
    row.update({"cluster_head_host": "all clusters pooled", "cluster_head_address": "", "n_epoch_pairs": n,
                "n_distinct_rho_values": len(set(pooled_x)), "n_pairs_rho_equal_1": sum(1 for x in pooled_x if x is not None and abs(x - 1.0) < 1e-9),
                "spearman_rho_e_vs_w_next": rho, "spearman_p": p,
                "spearman_rho_e_vs_w_next_over_W_raw_next": rho_c, "spearman_p_control": p_c,
                "spearman_R_k_vs_rho_k": rho_r, "spearman_p_R_vs_rho": p_r,
                "note": "pooled across clusters: w_k(e+1) mixes cluster sizes, so the pooled coefficient is dominated by the between-cluster ESG x activity differences, not by rho"})
    corr_rows.append(row)

    # summary of the Gini trajectory
    gs = [(r["epoch"], r["gini_delta_published_minus_input"]) for r in gini_rows if r["gini_delta_published_minus_input"] != ""]
    summ = OrderedDict(base)
    if len(gs) >= 2:
        xs = np.array([g[0] for g in gs], dtype=float)
        ys = np.array([g[1] for g in gs], dtype=float)
        slope = float(np.polyfit(xs, ys, 1)[0])
        summ.update({"n_epochs": len(gs), "gini_delta_first": ys[0], "gini_delta_last": ys[-1],
                     "gini_delta_max": float(ys.max()), "gini_delta_slope_per_epoch": slope,
                     "gini_published_last3_sd": float(np.std([r["gini_w_published"] for r in gini_rows[-3:] if r["gini_w_published"] != ""])) if len(gini_rows) >= 3 else "",
                     "verdict_descriptive": ("amplification by feedback (published Gini grows faster than input Gini)" if slope > 1e-4 else
                                             ("no amplification: published inequality mirrors the inputs" if abs(slope) <= 1e-4 else
                                              "published Gini grows slower than input Gini")),
                     "note": "descriptive slope of gini_published - gini_input over epochs; a plateau is consistent with lambda_w < 1 (Prop. 6.2)"})
    else:
        summ.update({"n_epochs": len(gs), "note": NC + "fewer than 2 epochs with published weights"})
    traffic = OrderedDict(base)
    traffic.update({"statement": "traffic generation is EXOGENOUS: tools/role_company.sh do_traffic() publishes one item every INTERVAL simulated seconds regardless of who mines; "
                                 "the only endogenous feedback channel is rho (reconciliation), as tested above"})
    return epoch_rows, gini_rows, corr_rows, [summ], [traffic]


# --------------------------------------------------------------------------
# notes and excel
# --------------------------------------------------------------------------

def methodology_notes(rd, base, sig_rows):
    notes = [
        ("phase separation", "phase 1 = raw rows traceable to log lines / txids; phase 2 = deterministic recomputation; phase 3 (this file) = the only place with tests, intervals and Monte Carlo"),
        ("model check", "delay recomputed from the logged score with the in-force weight map, W_tot and Phi reproduces the logged delay within %s s on every candidate row except those listed in phase2 manifest summary.rounds_delay_mismatch (weight-map view divergence at a confirm height)" % 0.0015),
        ("Wilson", "95%% Wilson score intervals for every proportion; no normal approximation of p_hat"),
        ("GoF", "chi-square asymptotic only when min E_i >= 5; otherwise exact multinomial by enumeration (<= 3 classes) or Monte Carlo (N=%d); gof_method is always declared; a round-by-round Monte Carlo respecting per-round shares is reported alongside" % MC_GOF),
        ("Beta(1,n)", "KS of G against 2*Dmax*Beta(1,n) tests the iid-uniform approximation of Prop. 5.18, which the thesis itself declares invalid for concentrated weights (D11); the KS against the exact Monte Carlo of the implemented sortition (same weights, W_tot, Phi) is the reference test"),
        ("Prop. 5.17", "score(2)-score(1) given argmin i* is Exp(W - w_i*); gaps are standardised by the rate of their round and tested against Exp(1)"),
        ("sigma", "; ".join("%s = %s (%s)" % (r["sigma_id"], _r(r["sigma_s"], 4), r["note"] or r["source"]) for r in sig_rows)),
        ("inversion bound", "the bound of Prop. 5.18 is evaluated with sigma_S2 (primary) and sigma_S1; the observed rate is compared with both and with a Gaussian-noise Monte Carlo at sigma_S2"),
        ("epoch and in-force weight", "epoch(h) = h // L + 1 (weight_engine.h, one function shared by both modules); a weight record confirmed at height h_c is in force for blocks h > h_c; the expected share of an epoch is the block-weighted mean of the in-force shares, never the weight stamped with that epoch"),
        ("Theta", "theta_code = sum of tau over the companies in clusters (WeightEngine::NetworkActivity, D1); theta_thesis = all non-coinbase transactions of the epoch's blocks (Def. 6.6); both reported, the fold uses theta_code"),
        ("integer weight", "the on-chain weight is round-half-away(w_k * kappa) clamped >= 1 (ToIntegerWeight, D2); recompute_match compares that integer"),
        ("malus points", "p(selfwrite) and p(badweight) are '[null]' in params.dat; the start-up line of the node declares 1 and 2 (D4)"),
        ("not computable on this campaign", "dump-function effect (none everywhere); malus trajectory and Delay-event recomputation (registry empty); VRF output y_i (not logged); per-event network latency and jitter (1 s log resolution, jitter 0 us); packet loss"),
        ("replication", "replication_id = 1 everywhere: single seed per run, no replicas"),
        ("algorithm-level test", "tools/valida_sortition_montecarlo.py (standalone Monte Carlo of Algorithm 1) is run at campaign level and copied to campaign/sortition_algorithm_validation.csv; it is a test of the formula, not of the C++ binary"),
        ("constants", "alpha = %.2f; MC GoF/KS N = %d; MC streak N = %d; RNG seed = %d" % (ALPHA, MC_GOF, MC_STREAK, SEED)),
    ]
    rows = []
    for k, v in notes:
        row = OrderedDict(base)
        row.update({"topic": k, "note": v})
        rows.append(row)
    return rows


def write_xlsx(path, header, sheets):
    """sheets: list of (sheet_name, [(block_title, rows)])."""
    if not USING_OPENPYXL:
        LOG.warning("openpyxl missing: %s not written (CSV twins are complete)", path)
        return False
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    for name, blocks in sheets:
        ws = wb.create_sheet(title=name[:31])
        ws.append(["parameters"] + list(header.keys()))
        ws.append([""] + [str(v) for v in header.values()])
        ws.cell(row=1, column=1).font = bold
        ws.append([])
        for title, rows in blocks:
            ws.append([title])
            ws.cell(row=ws.max_row, column=1).font = bold
            if not rows:
                ws.append(["(empty)"])
                ws.append([])
                continue
            cols = []
            for r in rows:
                for k in r:
                    if k not in cols:
                        cols.append(k)
            ws.append(cols)
            for c in range(1, len(cols) + 1):
                ws.cell(row=ws.max_row, column=c).font = bold
            for r in rows:
                ws.append([_r(r.get(k, ""), 6) if not isinstance(r.get(k, ""), (list, dict)) else str(r.get(k)) for k in cols])
            ws.append([])
        ws.freeze_panes = "A4"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return True


def analyze_run(p2, p3, p1):
    man, t = load_phase2(p2)
    rd = RunData(man, t)
    base = OrderedDict([("experiment_id", man["experiment_id"]), ("area", man["area"]),
                        ("replication_id", man["replication_id"]), ("target_block_time", man["target_block_time"]),
                        ("dumpfunction", man["dumpfunction"] or "none"), ("malus_enabled", man["malus_enabled"])])
    header = param_header(man)
    sig_rows, sigmas = sigma_rows(rd, base)
    val_rows, test_rows, live_rows, p517_rows = wpoa_epoch(rd, base, sigmas)
    gain_rows, fit_rows, lr_rows, _ = wpoa_longitudinal(rd, base)
    we_rows, gini_rows, corr_rows, gini_summ, traffic = weight_engine(rd, base)
    notes = methodology_notes(rd, base, sig_rows)
    config = C.read_csv_dicts(p1 / "config.csv")

    p3.mkdir(parents=True, exist_ok=True)
    tables = OrderedDict([
        ("wpoa_epoch_validators", val_rows), ("wpoa_epoch_tests", test_rows),
        ("wpoa_epoch_liveness_validators", live_rows), ("wpoa_prop517", p517_rows), ("wpoa_sigma", sig_rows),
        ("wpoa_longitudinal_validators", gain_rows), ("wpoa_longitudinal_fits", fit_rows),
        ("wpoa_longitudinal_logratios", lr_rows),
        ("weight_engine_epoch", we_rows), ("weight_engine_gini", gini_rows), ("weight_engine_correlations", corr_rows),
        ("weight_engine_gini_summary", gini_summ), ("weight_engine_traffic_model", traffic),
        ("config", config), ("methodology_notes", notes),
    ])
    counts = {k: C.write_csv(p3 / (k + ".csv"), _round_rows(v, 8)) for k, v in tables.items()}
    man_rows = [OrderedDict([("key", k), ("value", json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)])
                for k, v in man.items() if k not in ("epochs", "params", "startup_params")]
    man_rows += [OrderedDict([("key", "epochs"), ("value", json.dumps(man.get("epochs", []), ensure_ascii=False))])]
    xlsx = p3 / ("%s__%s.xlsx" % (man["experiment_id"], man["area"]))
    write_xlsx(xlsx, header, [
        ("wPoA_epoch", [("6.1.1-6.1.4 weights, p_i, Wilson 95%, errors (per epoch and pooled 'all')", val_rows),
                        ("6.1.2/6.1.5-6.1.9 GoF, timer race, concentration, streaks, liveness, time bar", test_rows),
                        ("6.1.8 MissedTurnRate per validator", live_rows),
                        ("Prop. 5.17 score gap per argmin class", p517_rows),
                        ("sigma estimates (S1 topology, S2 scheduler, S3 packet loss)", sig_rows)]),
        ("wPoA_longitudinal", [("6.2 gains and monotonicity (first vs last fully measured epoch)", gain_rows),
                               ("6.2 regressions", fit_rows),
                               ("6.2 pairwise log-ratios", lr_rows)]),
        ("WeightEngine_epoch", [("6.3.1 cluster table per epoch", we_rows),
                                ("6.3.3 Gini vs input baseline", gini_rows),
                                ("6.3.3 Gini trajectory summary", gini_summ),
                                ("6.3.2 lag-1 Spearman rho_k(e) -> w_k(e+1)", corr_rows),
                                ("traffic model", traffic)]),
        ("config", [("parameters with sources (phase 1)", config)]),
        ("manifest", [("phase-2 manifest", man_rows)]),
        ("notes", [("methodology notes and non-computable items", notes)]),
    ])
    summary = OrderedDict(list(header.items()))
    pooled = next((r for r in test_rows if r["epoch"] == "all"), {})
    fit1 = fit_rows[0] if fit_rows else {}
    corr_pool = corr_rows[-1] if corr_rows else {}
    gsum = gini_summ[0] if gini_summ else {}
    summary.update({
        "n_measured_blocks": man["summary"]["n_measured_blocks"], "n_rounds_with_scores": man["summary"]["n_rounds_with_scores"],
        "measured_epochs": len(rd.measured), "fully_measured_epochs": len(rd.fully_measured),
        "inversion_rate": pooled.get("inversion_rate", ""), "inversion_wilson95_low": pooled.get("inversion_wilson95_low", ""),
        "inversion_wilson95_high": pooled.get("inversion_wilson95_high", ""),
        "inversion_with_fork_n": pooled.get("inversion_with_fork_n", ""), "inversion_n": pooled.get("inversion_n", ""),
        "sigma_S1_s": sigmas[0] if sigmas[0] is not None else "", "sigma_S2_s": sigmas[1] if sigmas[1] is not None else "",
        "inversion_bound_S1": pooled.get("inversion_bound_sigma_S1", ""), "inversion_bound_S2": pooled.get("inversion_bound_sigma_S2", ""),
        "inversion_mc_prob_gaussian_S2": pooled.get("inversion_mc_prob_gaussian_sigma_S2", ""),
        "gof_pooled_method": pooled.get("gof_method", ""), "gof_pooled_p": pooled.get("gof_p_value", ""),
        "gof_pooled_min_expected": pooled.get("gof_min_expected", ""),
        "gof_pooled_mc_round_by_round_p": pooled.get("gof_mc_round_by_round_p", ""),
        "TV_pooled": pooled.get("TV", ""), "MaxAE_pooled": pooled.get("MaxAE_p", ""),
        "epochs_gof_reject_alpha05": sum(1 for r in test_rows if r["epoch"] != "all" and r["gof_reject_alpha05"] == 1),
        "epochs_gof_tested": sum(1 for r in test_rows if r["epoch"] != "all" and r["gof_p_value"] != ""),
        "ks_beta_pooled_p": pooled.get("ks_beta_p", ""), "ks_mc_pooled_p": pooled.get("ks_mc_p", ""),
        "G_median_over_Dmax": pooled.get("G_over_Dmax_median", ""),
        "HHI_p_hat_pooled": pooled.get("p_hat_HHI", ""), "HHI_p_theoretical_pooled": pooled.get("p_theoretical_HHI", ""),
        "nakamoto_1_2_p_hat": pooled.get("p_hat_nakamoto_1_2", ""),
        "L_max_observed": pooled.get("L_max_observed", ""), "L_max_mc_mean": pooled.get("L_max_mc_mean", ""),
        "L_max_mc_p": pooled.get("L_max_mc_p_value", ""),
        "BlockSuccessRate": pooled.get("BlockSuccessRate", ""), "n_forks": pooled.get("n_forks", ""),
        "InterBlockTime_median_s": pooled.get("InterBlockTime_median_s", ""), "InterBlockTime_p95_s": pooled.get("InterBlockTime_p95_s", ""),
        "time_bar_violations": pooled.get("time_bar_violations", ""), "n_rejects": pooled.get("n_rejects", ""),
        "delay_recompute_mismatch_rounds": pooled.get("delay_recompute_mismatch_rounds", ""),
        "phi_mismatch_rounds": pooled.get("phi_mismatch_rounds", ""),
        "glm_beta1_logw": fit1.get("beta1", ""), "glm_beta1_ci_low": fit1.get("beta1_ci_low", ""), "glm_beta1_ci_high": fit1.get("beta1_ci_high", ""),
        "weight_recompute_mismatches": man["summary"]["n_weight_recompute_mismatches"],
        "weight_rows": man["summary"]["n_weight_rows"],
        "gini_delta_last": gsum.get("gini_delta_last", ""), "gini_delta_slope": gsum.get("gini_delta_slope_per_epoch", ""),
        "gini_verdict": gsum.get("verdict_descriptive", ""),
        "spearman_rho_to_w_next_pooled": corr_pool.get("spearman_rho_e_vs_w_next", ""), "spearman_p_pooled": corr_pool.get("spearman_p", ""),
        "malus_registry_records": man["summary"]["malus_registry_records"],
        "xlsx": str(xlsx.relative_to(p3.parent.parent.parent)) if xlsx.exists() else "",
    })
    (p3 / "manifest.json").write_text(json.dumps(OrderedDict([("phase", 3), ("row_counts", counts), ("summary", summary),
                                                               ("constants", {"alpha": ALPHA, "mc_gof": MC_GOF, "mc_streak": MC_STREAK, "seed": SEED})]),
                                                  indent=1, ensure_ascii=False, default=str), encoding="utf-8")
    LOG.info("%s: phase3 written (%s)", man["run_id"], xlsx.name)
    return summary


# --------------------------------------------------------------------------
# campaign level
# --------------------------------------------------------------------------

def campaign(root, summaries, tools_dir):
    cdir = root / "campaign"
    cdir.mkdir(parents=True, exist_ok=True)
    rows = _round_rows(summaries, 6)
    C.write_csv(cdir / "campaign_summary.csv", rows)
    index = []
    for s in summaries:
        index.append({"run_id": "%s/%s" % (s["experiment_id"], s["area"]), "livello": s["area"], "stato": "completa",
                      "target_block_time": s["target_block_time"], "mining_diversity": s["mining_diversity"],
                      "weight_epoch_length": s["weight_epoch_length"], "weight_kappa": s["weight_kappa"],
                      "weight_lambda": s["weight_lambda"], "wpoa_sortition_delta": s["wpoa_sortition_delta"],
                      "wpoa_sortition_lambda": s["wpoa_sortition_lambda"], "wpoa_randao_lookback": s["wpoa_randao_lookback"],
                      "dump_function": s["dumpfunction"]})
    fams = C.build_families(index)
    fam_rows = [OrderedDict([("varied_parameter", f["parametro_variato"]), ("values", f["valori"]), ("n_runs", f["n_run"]),
                             ("run_ids", f["run_id"]), ("constants", f["costanti"])]) for f in fams]
    C.write_csv(cdir / "campaign_families.csv", fam_rows)
    manifest = OrderedDict([("phase", 3), ("n_runs", len(summaries)), ("hierarchy", "experiment (runN-tbt-Ts, sealed params) -> area -> epochs"),
                            ("experiments", OrderedDict())])
    for s in summaries:
        ex = manifest["experiments"].setdefault(s["experiment_id"], OrderedDict([
            ("target_block_time", s["target_block_time"]), ("weight_epoch_length", s["weight_epoch_length"]),
            ("dumpfunction", s["dumpfunction"]), ("malus_enabled", s["malus_enabled"]), ("areas", OrderedDict())]))
        ex["areas"][s["area"]] = {"replication_id": s["replication_id"], "measured_epochs": s["measured_epochs"],
                                  "n_measured_blocks": s["n_measured_blocks"], "xlsx": s["xlsx"]}
    # standalone algorithm-level Monte Carlo (reused as-is)
    mc_out = cdir / "sortition_algorithm_validation.csv"
    script = tools_dir / "valida_sortition_montecarlo.py"
    mc_note = ""
    if script.is_file():
        cmd = [sys.executable, str(script), "--estrazioni", str(MC_GOF), "--out", str(mc_out),
               "--peso-riconciliazione", str(tools_dir.parent / "analisi" / "esperimenti" / "*" / "*" / "peso_riconciliazione.csv")]
        try:
            res = subprocess.run(cmd, cwd=str(tools_dir.parent), capture_output=True, text=True, timeout=1800)
            (cdir / "sortition_algorithm_validation.log").write_text(res.stdout + "\n" + res.stderr, encoding="utf-8")
            mc_note = "exit %d" % res.returncode
        except (OSError, subprocess.TimeoutExpired) as exc:
            mc_note = "failed: %s" % exc
    else:
        mc_note = "script not found"
    manifest["sortition_algorithm_validation"] = {"script": str(script), "draws": MC_GOF, "status": mc_note,
                                                  "note": "test of the Efraimidis-Spirakis formula in Python, not of the C++ binary"}
    (cdir / "campaign_manifest.json").write_text(json.dumps(manifest, indent=1, ensure_ascii=False, default=str), encoding="utf-8")
    # excel: summary, families, side-by-side per family
    if USING_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "summary"
        cols = list(rows[0].keys()) if rows else []
        ws.append(cols)
        for r in rows:
            ws.append([r.get(k, "") for k in cols])
        ws.freeze_panes = "C2"
        ws = wb.create_sheet("families")
        ws.append(list(fam_rows[0].keys()) if fam_rows else ["(none)"])
        for r in fam_rows:
            ws.append(list(r.values()))
        ws = wb.create_sheet("by_family")
        by_id = {"%s/%s" % (r["experiment_id"], r["area"]): r for r in rows}
        metrics = [c for c in cols if c not in ("experiment_id", "area")]
        for f in fam_rows:
            ids = f["run_ids"].split()
            ws.append(["family: %s = {%s}" % (f["varied_parameter"], f["values"]), "constants: " + f["constants"]])
            ws.append(["metric"] + ids)
            for m in metrics:
                ws.append([m] + [by_id.get(i, {}).get(m, "") for i in ids])
            ws.append([])
        mc_rows = C.read_csv_dicts(mc_out)
        ws = wb.create_sheet("algorithm_mc")
        if mc_rows:
            ws.append(list(mc_rows[0].keys()))
            for r in mc_rows:
                ws.append(list(r.values()))
        else:
            ws.append(["not available: " + mc_note])
        wb.save(str(cdir / "campaign.xlsx"))
    LOG.info("campaign written in %s (%d runs, %d families, MC validator %s)", cdir, len(summaries), len(fam_rows), mc_note)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default="risultati")
    ap.add_argument("--only", default="")
    ap.add_argument("--campaign", action="store_true", help="also write the campaign-level files")
    ap.add_argument("--verbose", action="store_true")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)-7s %(message)s")
    root = Path(args.out).expanduser().resolve()
    summaries = []
    for p2 in sorted(root.glob("*/*/phase2")):
        rid = "%s/%s" % (p2.parent.parent.name, p2.parent.name)
        if args.only not in rid:
            continue
        try:
            summaries.append(analyze_run(p2, p2.parent / "phase3", p2.parent / "phase1"))
        except Exception as exc:
            LOG.exception("%s: phase3 failed: %s", rid, exc)
    if not summaries:
        LOG.error("no phase2 directories under %s", root)
        return 1
    if args.campaign:
        campaign(root, summaries, Path(__file__).resolve().parent.parent)
    return 0


if __name__ == "__main__":
    sys.exit(main())
